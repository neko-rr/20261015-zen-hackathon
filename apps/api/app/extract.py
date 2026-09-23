"""手元資料の形式判定と索引用テキスト／画像派生。"""

from __future__ import annotations

import io
import logging
from pathlib import Path

logger = logging.getLogger(__name__)

MAX_MATERIAL_BYTES = 100 * 1024 * 1024
MAX_MATERIALS_COUNT = 30
MAX_MATERIALS_TOTAL_BYTES = 2 * 1024 * 1024 * 1024
INDEX_IMAGE_MAX_EDGE = 2048

MATERIAL_KINDS = frozenset({"own_work", "third_party"})

# MIME または拡張子 → media_type
_MIME_TO_MEDIA: dict[str, str] = {
    "image/jpeg": "image",
    "image/png": "image",
    "image/webp": "image",
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "word",
    "application/msword": "word",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "excel",
    "application/vnd.ms-excel": "excel",
    "text/plain": "note",
    "text/markdown": "note",
}

_EXT_TO_MEDIA: dict[str, str] = {
    ".jpg": "image",
    ".jpeg": "image",
    ".png": "image",
    ".webp": "image",
    ".pdf": "pdf",
    ".docx": "word",
    ".doc": "word",
    ".xlsx": "excel",
    ".xls": "excel",
    ".md": "note",
    ".txt": "note",
}

_EXT_TO_MIME: dict[str, str] = {
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".webp": "image/webp",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".doc": "application/msword",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xls": "application/vnd.ms-excel",
    ".md": "text/markdown",
    ".txt": "text/plain",
}


def resolve_media_type(filename: str, content_type: str | None) -> str | None:
    mime = (content_type or "").split(";")[0].strip().lower()
    if mime in _MIME_TO_MEDIA:
        return _MIME_TO_MEDIA[mime]
    ext = Path(filename or "").suffix.lower()
    return _EXT_TO_MEDIA.get(ext)


def guess_mime(filename: str, content_type: str | None) -> str:
    mime = (content_type or "").split(";")[0].strip().lower()
    if mime and mime != "application/octet-stream":
        return mime
    ext = Path(filename or "").suffix.lower()
    return _EXT_TO_MIME.get(ext, "application/octet-stream")


def prepare_index_payload(
    media_type: str,
    filename: str,
    raw: bytes,
) -> tuple[str, bytes | None, str]:
    """索引用テキストと、画像なら縮小バイトを返す。"""
    title = Path(filename or "untitled").name
    if media_type == "image":
        derivative = _shrink_image(raw)
        text = f"{title}\nimage material"
        return text, derivative, "unknown"
    if media_type == "note":
        try:
            body = raw.decode("utf-8")
        except UnicodeDecodeError:
            body = raw.decode("utf-8", errors="replace")
        return f"{title}\n{body}"[:50_000], None, "unknown"
    if media_type == "word":
        return f"{title}\n{_extract_docx(raw)}"[:50_000], None, "unknown"
    if media_type == "excel":
        return f"{title}\n{_extract_xlsx(raw)}"[:50_000], None, "unknown"
    if media_type == "pdf":
        body, page = _extract_pdf(raw)
        return f"{title}\n{body}"[:50_000], None, page
    return title, None, "unknown"


def _shrink_image(raw: bytes) -> bytes | None:
    try:
        from PIL import Image
    except ImportError:
        logger.exception("pillow missing")
        return None
    try:
        image = Image.open(io.BytesIO(raw))
        image = image.convert("RGB")
        image.thumbnail((INDEX_IMAGE_MAX_EDGE, INDEX_IMAGE_MAX_EDGE))
        out = io.BytesIO()
        image.save(out, format="JPEG", quality=85)
        return out.getvalue()
    except Exception:
        logger.exception("image shrink failed")
        return None


def _extract_docx(raw: bytes) -> str:
    try:
        from docx import Document
    except ImportError:
        logger.exception("python-docx missing")
        return ""
    try:
        document = Document(io.BytesIO(raw))
        return "\n".join(p.text for p in document.paragraphs if p.text.strip())
    except Exception:
        logger.exception("docx extract failed")
        return ""


def _extract_xlsx(raw: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        logger.exception("openpyxl missing")
        return ""
    try:
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        lines: list[str] = []
        for sheet in book.worksheets:
            lines.append(f"# {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(cell) for cell in row if cell is not None]
                if cells:
                    lines.append("\t".join(cells))
        return "\n".join(lines)
    except Exception:
        logger.exception("xlsx extract failed")
        return ""


def _extract_pdf(raw: bytes) -> tuple[str, str]:
    try:
        from pypdf import PdfReader
    except ImportError:
        logger.exception("pypdf missing")
        return "", "unknown"
    try:
        reader = PdfReader(io.BytesIO(raw))
        parts: list[str] = []
        for index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                parts.append(f"[page {index}]\n{text}")
        page_label = str(len(reader.pages)) if reader.pages else "unknown"
        return "\n".join(parts), page_label
    except Exception:
        logger.exception("pdf extract failed")
        return "", "unknown"
